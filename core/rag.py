"""Local RAG — nomic-embed-text via Ollama + FAISS vector index."""
import os, json, threading, requests
from typing import List, Dict, Optional, Tuple
from core.logger import get_logger
from config import DB_PATH, get
log = get_logger("rag")

try: import numpy as np; NP_OK = True
except ImportError: NP_OK = False
try: import faiss; FAISS_OK = True
except ImportError: FAISS_OK = False; log.warning("faiss-cpu not installed: pip install faiss-cpu")

EMBED_MODEL   = "nomic-embed-text"
EMBED_DIM     = 768
CHUNK_SIZE    = 500
CHUNK_OVERLAP = 50
_DIR          = os.path.dirname(DB_PATH)
INDEX_PATH    = os.path.join(_DIR, "rag.index")
META_PATH     = os.path.join(_DIR, "rag_meta.json")
SUPPORTED     = {".txt",".md",".py",".js",".ts",".json",".yaml",".yml",".csv",
                 ".log",".html",".sql",".sh",".bat",".pdf",".docx",".xlsx"}

def _chunk(text, size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    chunks=[]; start=0
    while start < len(text):
        chunks.append(text[start:start+size]); start += size-overlap
    return [c.strip() for c in chunks if c.strip()]

def _read(path):
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(path) as p: return "\n".join(pg.extract_text() or "" for pg in p.pages)
        elif ext == ".docx":
            from docx import Document
            return "\n".join(p.text for p in Document(path).paragraphs)
        elif ext in (".xlsx",".xls"):
            import openpyxl; wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
            rows=[]; [rows.extend(" ".join(str(c) for c in r if c) for r in ws.iter_rows(values_only=True)) for ws in wb.worksheets]
            wb.close(); return "\n".join(rows)
        else:
            with open(path,"r",encoding="utf-8",errors="replace") as f: return f.read()
    except Exception as e: log.warning(f"RAG read {path}: {e}"); return ""

def _embed(texts):
    if not NP_OK: return None
    vecs=[]
    for text in texts:
        try:
            r=requests.post(f"{get('ollama_base_url')}/api/embeddings",
                json={"model":EMBED_MODEL,"prompt":text},timeout=30)
            r.raise_for_status(); vec=r.json().get("embedding",[])
            if vec: vecs.append(vec)
        except Exception as e: log.error(f"Embed: {e}"); return None
    if not vecs: return None
    arr=np.array(vecs,dtype=np.float32)
    norms=np.linalg.norm(arr,axis=1,keepdims=True)
    return arr/(norms+1e-8)

class RAGEngine:
    def __init__(self):
        self._index=None; self._meta=[]; self._lock=threading.Lock(); self._load()

    def _load(self):
        if not FAISS_OK or not NP_OK: return
        try:
            if os.path.exists(INDEX_PATH) and os.path.exists(META_PATH):
                self._index=faiss.read_index(INDEX_PATH)
                with open(META_PATH) as f: self._meta=json.load(f)
                log.info(f"RAG: {len(self._meta)} chunks loaded")
            else:
                self._index=faiss.IndexFlatIP(EMBED_DIM)
        except Exception as e: log.error(f"RAG load: {e}"); self._index=faiss.IndexFlatIP(EMBED_DIM) if FAISS_OK else None

    def _save(self):
        try:
            if self._index: faiss.write_index(self._index, INDEX_PATH)
            with open(META_PATH,"w") as f: json.dump(self._meta, f)
        except Exception as e: log.error(f"RAG save: {e}")

    @property
    def chunk_count(self): return len(self._meta)

    def index_directory(self, path, on_progress=None):
        if not FAISS_OK: return 0,0
        path=os.path.expanduser(path)
        if not os.path.exists(path): return 0,0
        files=[]
        for root,_,fnames in os.walk(path):
            if any(p.startswith(".") for p in root.split(os.sep)): continue
            for fn in fnames:
                if os.path.splitext(fn)[1].lower() in SUPPORTED: files.append(os.path.join(root,fn))
        done=0; total_chunks=0
        for i,fpath in enumerate(files):
            if on_progress: on_progress(i+1,len(files),fpath)
            n=self.index_file(fpath)
            if n>0: done+=1; total_chunks+=n
        self._save(); log.info(f"RAG: {done} files, {total_chunks} chunks"); return done,total_chunks

    def index_file(self, path):
        if not FAISS_OK: return 0
        path=os.path.expanduser(path)
        if not os.path.exists(path) or os.path.splitext(path)[1].lower() not in SUPPORTED: return 0
        mtime=str(os.path.getmtime(path))
        if any(m.get("path")==path and m.get("mtime")==mtime for m in self._meta): return 0
        text=_read(path)
        if not text.strip(): return 0
        chunks=_chunk(text)
        if not chunks: return 0
        vecs=_embed(chunks)
        if vecs is None or len(vecs)!=len(chunks): return 0
        with self._lock:
            self._index.add(vecs)
            fname=os.path.basename(path)
            for idx,chunk in enumerate(chunks):
                self._meta.append({"path":path,"filename":fname,"chunk_idx":idx,"chunk":chunk,"mtime":mtime})
        return len(chunks)

    def search(self, query, top_k=5):
        if not FAISS_OK or not self._meta: return []
        vec=_embed([query])
        if vec is None: return []
        k=min(top_k,len(self._meta))
        scores,indices=self._index.search(vec,k)
        results=[]
        for score,idx in zip(scores[0],indices[0]):
            if 0<=idx<len(self._meta):
                m=self._meta[idx]
                results.append({"filename":m["filename"],"path":m["path"],"chunk":m["chunk"],"score":float(score)})
        return results

    def format_results(self, results):
        if not results: return "No relevant documents found."
        return "\n\n---\n\n".join(f"[{i+1}] {r['filename']} (score:{r['score']:.2f})\n{r['chunk'][:400]}"
                                   for i,r in enumerate(results))

    def get_rag_tool_description(self):
        return f"\n[RAG: query]  → Search your {self.chunk_count} indexed document chunks\n"
